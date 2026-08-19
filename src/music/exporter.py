"""
======================================================================
  导出模块 - raw(原始英文) + translated(中文汉化) 双目录输出
======================================================================
"""

from __future__ import annotations

import json
import os
from datetime import datetime

from .models import SongInfo, AlbumInfo, SearchResult


# ================================================================
#  主入口
# ================================================================
def export_all_formats(result: SearchResult, base_dir: str = "music_data"):
    if not result.songs:
        print("  无歌曲数据，跳过导出")
        return {}

    first = result.songs[0]
    singer = first.singer or "未知"
    kw = result.keyword.replace(singer, "").strip() or first.song_name
    dir_name = f"{singer} {kw}"
    safe_dir = _safe(dir_name)

    # ---- raw 目录（原始英文） ----
    raw_root = os.path.join(base_dir, safe_dir, "raw")
    os.makedirs(raw_root, exist_ok=True)
    _export_raw_json(result, os.path.join(raw_root, "full_raw.json"))

    # ---- translated 目录（中文汉化） ----
    trans_root = os.path.join(base_dir, safe_dir, "translated")
    os.makedirs(trans_root, exist_ok=True)

    results = {}

    # 目录结构
    sp = _export_directory_structure(result, trans_root)
    results["汇总JSON"] = sp

    # 完整JSON
    fp = _export_full_json(result, os.path.join(trans_root, f"{safe_dir}_full.json"))
    results["完整JSON"] = fp

    # Excel
    xp = _export_excel(result, os.path.join(trans_root, f"{safe_dir}.xlsx"))
    results["Excel"] = xp

    # Markdown
    mp = _export_markdown(result, os.path.join(trans_root, f"{safe_dir}.md"))
    results["Markdown"] = mp

    results["原始JSON"] = os.path.join(raw_root, "full_raw.json")
    results["输出目录"] = os.path.join(base_dir, safe_dir)

    return results


# ================================================================
#  Raw JSON（原始 Pydantic 模型字段名）
# ================================================================
def _export_raw_json(result: SearchResult, filepath: str):
    data = result.model_dump(exclude_defaults=False)
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2, default=str)


# ================================================================
#  目录结构
# ================================================================
def _export_directory_structure(result: SearchResult, root: str) -> str:
    all_platforms_str = "、".join(result.platforms_available) if result.platforms_available else ""
    summary_songs = []
    platform_count = {}
    seq = {}

    for s in result.songs:
        plat = s.platform_name or s.platform or "未知"
        seq[plat] = seq.get(plat, 0) + 1

        record = _build_full_record(s, all_platforms_str)
        summary_songs.append(record)
        platform_count[plat] = platform_count.get(plat, 0) + 1

        plat_dir = os.path.join(root, plat)
        os.makedirs(plat_dir, exist_ok=True)
        fname = f"{seq[plat]:02d}_{_safe(s.song_name)}_{_safe(s.singer or '未知')}.json"
        with open(os.path.join(plat_dir, fname), "w", encoding="utf-8") as f:
            json.dump(record, f, ensure_ascii=False, indent=2)

    summary = {
        "采集时间": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "总歌曲数": len(summary_songs),
        "各平台统计": platform_count,
        "所有歌曲": summary_songs,
    }
    sp = os.path.join(root, "all_platforms_summary.json")
    with open(sp, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
    return sp


def _build_full_record(s: SongInfo, all_platforms_str: str) -> dict:
    record = {
        "歌曲ID": s.song_id,
        "歌曲名": s.song_name or "",
        "完整标题": s.title or s.song_name or "",
        "副标题": s.subtitle or "",
        "别名": s.alias or [],
        "歌手": s.singer or "",
        "歌手ID": s.singer_id or "",
        "歌手MID": s.singer_mid or "",
        "语种": s.language or "",
        "歌词": s.lyrics or "",
        "歌曲图片": s.cover_url or "",
        "专辑名": s.album_name or "",
        "专辑ID": s.album_id or "",
        "曲目序号": s.track_no,
        "碟片号": s.disc_no or "",
        "时长(秒)": s.duration,
        "时长(毫秒)": s.interval,
        "BPM": s.bpm,
        "流派": s.genre or "",
        "发行时间": s.release_date or "",
        "发布时间戳": s.publish_time,
        "热度": s.popularity,
        "MV_ID": s.mv.id if s.mv else "",
        "MV_VID": s.mv.vid if s.mv else "",
        "付费类型": s.pay.fee if s.pay else 0,
        "付费试听": s.pay.pay_play if s.pay else 0,
        "付费下载": s.pay.pay_down if s.pay else 0,
        "单曲价格": s.pay.price_track if s.pay else 0,
        "音质文件": [
            {"音质": f.quality, "大小(bytes)": f.size, "扩展名": f.extension,
             "比特率": f.bitrate, "采样率": f.sr, "直链": f.url}
            for f in (s.files or [])
        ],
        "文件大小_128k": s.size128,
        "文件大小_320k": s.size320,
        "文件大小_FLAC": s.sizeflac,
        "URL链接": s.url or "",
        "音乐直链": s.play_url or "",
        "平台": s.platform_name or s.platform or "",
        "平台标识": s.platform or "",
        "上线平台": all_platforms_str or "",
        "模式": s.mode or "",
        "采集时间": s.collected_at,
        "专辑内歌曲": _album_song_names(s),
        "专辑简介": _album_desc(s),
        "专辑类型": _album_type(s),
        "API原始数据": s.extra if s.extra else {},
        "API数据来源平台": s.platform_name or s.platform or "",
    }
    return record


def _album_song_names(s: SongInfo) -> list:
    album = (s.extra or {}).get("_album", {})
    if album:
        songs = album.get("songs", [])
        result = []
        for asong in songs:
            name = asong.get("song_name") or asong.get("name") or ""
            url = asong.get("url", "")
            if url:
                result.append({"歌曲名": name, "链接": url})
            else:
                result.append(name)
        return result
    return []

def _album_desc(s: SongInfo) -> str:
    album = (s.extra or {}).get("_album", {})
    return album.get("description", album.get("desc", "")) or ""

def _album_type(s: SongInfo) -> str:
    album = (s.extra or {}).get("_album", {})
    return album.get("album_type", album.get("type", "")) or ""


# ================================================================
#  完整JSON（全中文）
# ================================================================
def _export_full_json(result: SearchResult, filepath: str):
    data = {
        "检索关键词": result.keyword,
        "上线平台": "、".join(result.platforms_available) if result.platforms_available else "",
        "采集时间": result.collected_at,
        "总歌曲数": result.total_songs,
        "总专辑数": result.total_albums,
        "歌曲列表": [_song_to_chinese(s, "、".join(result.platforms_available)) for s in result.songs],
        "专辑列表": [_album_to_chinese(a) for a in (result.albums or [])],
    }
    if result.fivesing_results:
        data["5sing分模式"] = {
            mode: [_song_to_chinese(s, "") for s in songs]
            for mode, songs in result.fivesing_results.items()
        }
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2, default=str)
    return filepath


def _song_to_chinese(s: SongInfo, all_platforms_str: str = "") -> dict:
    return {
        "歌曲ID": s.song_id,
        "歌曲名": s.song_name or "",
        "完整标题": s.title or s.song_name or "",
        "副标题": s.subtitle or "",
        "别名": s.alias or [],
        "歌手": s.singer or "",
        "歌手ID": s.singer_id or "",
        "歌手MID": s.singer_mid or "",
        "语种": s.language or "",
        "歌词": s.lyrics or "",
        "歌曲图片": s.cover_url or "",
        "专辑名": s.album_name or "",
        "专辑ID": s.album_id or "",
        "曲目序号": s.track_no,
        "碟片号": s.disc_no or "",
        "时长(秒)": s.duration,
        "时长(毫秒)": s.interval,
        "BPM": s.bpm,
        "流派": s.genre or "",
        "发行时间": s.release_date or "",
        "发布时间戳": s.publish_time,
        "热度": s.popularity,
        "MV_ID": s.mv.id if s.mv else "",
        "MV_VID": s.mv.vid if s.mv else "",
        "付费类型": s.pay.fee if s.pay else 0,
        "付费试听": s.pay.pay_play if s.pay else 0,
        "付费下载": s.pay.pay_down if s.pay else 0,
        "单曲价格": s.pay.price_track if s.pay else 0,
        "音质文件": [
            {"音质": f.quality, "大小(bytes)": f.size, "扩展名": f.extension,
             "比特率": f.bitrate, "采样率": f.sr, "直链": f.url}
            for f in (s.files or [])
        ],
        "文件大小_128k": s.size128,
        "文件大小_320k": s.size320,
        "文件大小_FLAC": s.sizeflac,
        "URL链接": s.url or "",
        "音乐直链": s.play_url or "",
        "平台": s.platform_name or s.platform or "",
        "平台标识": s.platform or "",
        "上线平台": all_platforms_str or "",
        "模式": s.mode or "",
        "采集时间": s.collected_at,
        "专辑内歌曲": _album_song_names(s),
        "专辑简介": _album_desc(s),
        "专辑类型": _album_type(s),
        "API原始数据": s.extra if s.extra else {},
        "API数据来源平台": s.platform_name or s.platform or "",
    }


def _album_to_chinese(a: AlbumInfo) -> dict:
    return {
        "专辑ID": a.album_id or "",
        "专辑名": a.album_name or "",
        "完整标题": a.title or a.album_name or "",
        "副标题": a.subtitle or "",
        "歌手": a.singer or "",
        "歌手ID": a.singer_id or "",
        "发行时间": a.release_date or "",
        "发布时间戳": a.publish_time,
        "专辑简介": a.description or "",
        "专辑类型": a.album_type or "",
        "语种": a.language or "",
        "流派": a.genre or "",
        "发行公司": a.company or "",
        "封面URL": a.cover_url or "",
        "曲目数": a.song_count,
        "曲目列表": [_song_to_chinese(s, "") for s in (a.songs or [])],
        "平台": a.platform_name or a.platform or "",
        "平台标识": a.platform or "",
        "URL链接": a.url or "",
        "采集时间": a.collected_at,
        "API原始数据": a.extra if a.extra else {},
        "API数据来源平台": a.platform_name or a.platform or "",
    }


# ================================================================
#  Excel
# ================================================================
def _export_excel(result: SearchResult, filepath: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    import json as _json

    wb = Workbook()
    hf = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
    hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    hfill2 = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
    ha = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ca = Alignment(vertical="top", wrap_text=True)
    bd = Border(left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"), bottom=Side(style="thin"))

    def hdr(ws, headers, fill=None):
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font, cell.fill, cell.alignment, cell.border = hf, (fill or hfill), ha, bd
    def row(ws, r, vals):
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment, cell.border = ca, bd
    def wid(ws, mn=8, mx=60):
        for col in ws.columns:
            lens = [sum(2 if ord(ch) > 127 else 1 for ch in str(c.value or "")[:100]) for c in col]
            ws.column_dimensions[col[0].column_letter].width = min(max((max(lens) if lens else mn) + 2, mn), mx)

    # Sheet 1
    ws1 = wb.active
    ws1.title = "概览"
    for i, (k, v) in enumerate([("关键词", result.keyword), ("平台", "、".join(result.platforms_available)),
                                 ("歌曲数", result.total_songs), ("专辑数", result.total_albums),
                                 ("采集时间", result.collected_at)], 1):
        ws1.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws1.cell(row=i, column=2, value=v)
    wid(ws1)

    # Sheet 2
    ws2 = wb.create_sheet("歌曲完整信息")
    cols = ["序号","平台","歌曲ID","歌曲名","歌手","语种","时长(s)","BPM","流派",
            "专辑名","发行时间","热度","MV_ID","歌词字数","音乐直链","URL","API原始数据"]
    hdr(ws2, cols)
    for i, s in enumerate(result.songs, 1):
        row(ws2, i + 1, [i, s.platform_name, s.song_id, s.song_name, s.singer, s.language or "?",
              s.duration, s.bpm or "", s.genre or "", s.album_name or "", s.release_date or "",
              s.popularity, s.mv.id if s.mv else "", len(s.lyrics), s.play_url, s.url,
              _json.dumps(s.extra, ensure_ascii=False, indent=1) if s.extra else ""])
    wid(ws2, mx=80)
    ws2.auto_filter.ref = ws2.dimensions
    ws2.freeze_panes = "A2"

    # Sheet 3: API明细
    ws3 = wb.create_sheet("API原始数据明细")
    hdr(ws3, ["歌曲序号","平台","歌曲名","API字段名","字段值"], hfill2)
    ri = 2
    for i, s in enumerate(result.songs, 1):
        if s.extra:
            for key, val in s.extra.items():
                val_str = _json.dumps(val, ensure_ascii=False) if not isinstance(val, str) else str(val)
                row(ws3, ri, [i, s.platform_name, s.song_name, key, val_str[:500]])
                ri += 1
    ws3.column_dimensions["D"].width = 25
    ws3.column_dimensions["E"].width = 80
    ws3.auto_filter.ref = ws3.dimensions
    ws3.freeze_panes = "A2"

    # Sheet 4: 歌词
    ws4 = wb.create_sheet("歌词")
    hdr(ws4, ["序号","平台","歌曲名","歌手","完整歌词"])
    ri = 2
    for i, s in enumerate(result.songs, 1):
        if s.lyrics:
            row(ws4, ri, [i, s.platform_name, s.song_name, s.singer, s.lyrics])
            ri += 1
    ws4.column_dimensions["E"].width = 80

    # Sheet 5: 专辑
    if result.albums:
        ws5 = wb.create_sheet("专辑信息")
        hdr(ws5, ["序号","平台","专辑ID","专辑名","歌手","类型","语种","发行日期","曲目数","封面","简介","API原始数据"])
        for i, a in enumerate(result.albums, 1):
            row(ws5, i + 1, [i, a.platform_name, a.album_id, a.album_name, a.singer,
                  a.album_type or "?", a.language or "?", a.release_date, a.song_count,
                  a.cover_url, (a.description or "")[:300],
                  _json.dumps(a.extra, ensure_ascii=False, indent=1) if a.extra else ""])
        wid(ws5, mx=80)

    wb.save(filepath)
    return filepath


# ================================================================
#  Markdown
# ================================================================
def _export_markdown(result: SearchResult, filepath: str):
    import json as _json
    lines = []
    a = lines.append
    a("# 音乐信息获取报告")
    a("")
    a(f"- **关键词**: `{result.keyword}`")
    a(f"- **平台**: {', '.join(result.platforms_available) or '无'}")
    a(f"- **歌曲**: {result.total_songs} 首 | **专辑**: {result.total_albums} 个")
    a(f"- **采集时间**: {result.collected_at}")
    a("")
    a("---")
    a("")
    a("## 歌曲总表")
    a("")
    a("| # | 平台 | 歌曲名 | 歌手 | 时长 | 专辑 | 歌词 |")
    a("|---|------|--------|------|------|------|------|")
    for i, s in enumerate(result.songs, 1):
        dur = f"{s.duration // 60}:{s.duration % 60:02d}" if s.duration else "?"
        a(f"| {i} | {s.platform_name} | {_esc(s.song_name)} | {_esc(s.singer)} | {dur} | {_esc(s.album_name) or '-'} | {len(s.lyrics)}字 |")
    a("")

    for i, s in enumerate(result.songs, 1):
        a(f"## [{i}] {s.platform_name} - {s.song_name}")
        a("")
        for k, v in [
            ("歌曲ID", s.song_id), ("歌手", s.singer), ("语种", s.language or "?"),
            ("时长", f"{s.duration}s"), ("BPM", s.bpm), ("流派", s.genre or "?"),
            ("专辑", s.album_name or "-"), ("发行", s.release_date or "?"),
            ("热度", s.popularity), ("MV_ID", s.mv.id if s.mv else ""),
            ("封面", s.cover_url or "-"), ("直链", s.play_url or "-"),
            ("链接", s.url), ("平台标识", s.platform),
        ]:
            a(f"- **{k}**: {str(v).replace('|','\\\\|')}")

        if s.files:
            a("")
            a("| 音质 | 大小(KB) | 扩展名 | 比特率 |")
            a("|------|----------|--------|--------|")
            for f in s.files:
                a(f"| {f.quality} | {f.size//1024} | {f.extension} | {f.bitrate} |")

        album_inner = _album_song_names(s)
        if album_inner:
            a("")
            a("**专辑内歌曲**:")
            for ai in album_inner:
                if isinstance(ai, dict):
                    a(f"- [{ai.get('歌曲名','')}]({ai.get('链接','')})")
                else:
                    a(f"- {ai}")

        if s.lyrics:
            a("")
            a(f"<details><summary>歌词 ({len(s.lyrics)}字)</summary>")
            a("```")
            a(s.lyrics[:3000])
            a("```")
            a("</details>")

        if s.extra:
            a("")
            a(f"<details><summary>API原始数据 [{s.platform_name}] ({len(s.extra)}字段)</summary>")
            a("```json")
            a(_json.dumps(s.extra, ensure_ascii=False, indent=2)[:5000])
            a("```")
            a("</details>")

        a("")
        a("---")
        a("")

    if result.albums:
        a("## 专辑列表")
        for alb in result.albums:
            a(f"### [{alb.platform_name}] {alb.album_name}")
            for k, v in [("歌手", alb.singer), ("类型", alb.album_type), ("语种", alb.language),
                         ("日期", alb.release_date), ("曲目", alb.song_count), ("公司", alb.company), ("链接", alb.url)]:
                a(f"- **{k}**: {v}")
            a("")

    with open(filepath, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return filepath


# ================================================================
def _safe(name: str) -> str:
    for ch in '/\\:*?"<>|':
        name = name.replace(ch, "-")
    return name or "unknown"

def _esc(text: str) -> str:
    return (text or "").replace("|", "\\|").replace("\n", " ")
