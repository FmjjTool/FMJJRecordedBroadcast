"""音乐账号追踪：单账号单平台配置 + 采集 + xlsx 长期缓存。

每个账号：{name: 昵称, url: 链接, platform: netease/qq/kugou/kuwo/fivesing,
          kind: user/artist/singer（可空，采集时从链接识别）}。
采集结果追加保存到 导出目录/音乐账号数据/<昵称>.xlsx。
"""

import json
import re
import sys
from datetime import datetime
from pathlib import Path

BASE_DIR = Path(__file__).parent

PLATFORM_NAMES = {"netease": "网易云", "qq": "QQ音乐", "kugou": "酷狗",
                  "kuwo": "酷我", "fivesing": "5sing"}

# 组号规范（与「账号追踪」一致）：必须为重复数字，如 11、22、33 ... 99
GROUP_ID_PATTERN = re.compile(r"^(\d)\1$")

# xlsx 缓存字段（与采集返回的 record 字段对齐）
_RECORD_FIELDS = [
    "platform", "kind", "name", "fans", "following", "listens",
    "songs", "albums", "mvs", "works", "sign", "region",
    "birthday", "level", "extra", "avatar", "error",
]

_NUM_FIELDS = {"fans", "following", "listens", "songs", "albums", "mvs", "works", "level"}


def _data_dir() -> Path:
    """账号/分组数据目录：源码运行=模块旁；打包后=exe 旁 config/。"""
    if getattr(sys, "frozen", False):
        from toolbox.config_manager import get_config_dir
        return get_config_dir()
    return BASE_DIR


ACCOUNTS_FILE = _data_dir() / "music_accounts.json"
GROUPS_FILE = _data_dir() / "music_groups.json"


def _norm_account(a):
    """补全账号字段默认值，保证加载后的账号结构统一。"""
    a = dict(a)
    a.setdefault("group_id", "")
    a.setdefault("kind", "")
    return a


def load_accounts():
    """加载追踪账号列表：[{name, url, platform, kind, group_id}]。"""
    if ACCOUNTS_FILE.exists():
        try:
            with open(ACCOUNTS_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
            if isinstance(data, list):
                return [_norm_account(a) for a in data if isinstance(a, dict)]
        except Exception:
            pass
    return []


def save_accounts(accounts):
    with open(ACCOUNTS_FILE, "w", encoding="utf-8") as f:
        json.dump(accounts, f, ensure_ascii=False, indent=2)


def move_account(index: int, delta: int) -> bool:
    """上移/下移账号（index 为当前行号，delta=-1/+1），调整后立即保存。"""
    accounts = load_accounts()
    new = index + delta
    if not (0 <= index < len(accounts)) or not (0 <= new < len(accounts)):
        return False
    accounts[index], accounts[new] = accounts[new], accounts[index]
    save_accounts(accounts)
    return True


# ==================== 分组 ====================

def get_groups():
    """返回 {组号: 组名} 映射（读 music_groups.json）。"""
    if GROUPS_FILE.exists():
        try:
            with open(GROUPS_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
            if isinstance(data, dict):
                return {str(k): str(v) for k, v in data.items() if str(k)}
        except Exception:
            pass
    return {}


def save_group(group_id, group_name):
    group_id = (group_id or "").strip()
    group_name = (group_name or "").strip()
    if not group_id or not group_name:
        return False
    groups = get_groups()
    groups[group_id] = group_name
    try:
        with open(GROUPS_FILE, "w", encoding="utf-8") as f:
            json.dump(groups, f, ensure_ascii=False, indent=2)
        return True
    except Exception:
        return False


def delete_group(group_id):
    groups = get_groups()
    if group_id in groups:
        groups.pop(group_id)
        with open(GROUPS_FILE, "w", encoding="utf-8") as f:
            json.dump(groups, f, ensure_ascii=False, indent=2)
    accounts = load_accounts()
    changed = False
    for a in accounts:
        if a.get("group_id") == group_id:
            a["group_id"] = ""
            changed = True
    if changed:
        save_accounts(accounts)


def move_group(group_id, delta):
    """上移/下移分组（重排 music_groups.json 键顺序，列表/下拉显示顺序随之变化）。"""
    groups = get_groups()
    keys = list(groups.keys())
    if group_id not in keys:
        return False
    idx = keys.index(group_id)
    j = idx + (1 if delta > 0 else -1)
    if not (0 <= j < len(keys)):
        return False
    keys[idx], keys[j] = keys[j], keys[idx]
    reordered = {k: groups[k] for k in keys}
    try:
        with open(GROUPS_FILE, "w", encoding="utf-8") as f:
            json.dump(reordered, f, ensure_ascii=False, indent=2)
        return True
    except Exception:
        return False


def get_accounts_by_group(group_id):
    return [a for a in load_accounts() if a.get("group_id") == group_id]


# ==================== 链接解析 ====================

# (平台, 正则, 类型)——顺序即优先级
_URL_PATTERNS = [
    ("netease", r'music\.163\.com/(?:#/)?user/home\?id=(\d+)', "user"),
    ("netease", r'music\.163\.com/(?:#/)?home\?id=(\d+)', "user"),
    ("netease", r'music\.163\.com/(?:#/)?artist\?id=(\d+)', "artist"),
    ("qq", r'y\.qq\.com/n/ryqq/singer/([A-Za-z0-9]+)', "singer"),
    ("kugou", r'kugou\.com/yy/singer/home/(\d+)', "singer"),
    ("kugou", r'm\.kugou\.com/singer/(\d+)', "singer"),
    ("kugou", r'kugou\.com/singer/(\d+)', "singer"),
    ("kuwo", r'kuwo\.cn/singer_detail/(\d+)', "singer"),
    ("fivesing", r'5sing\.kugou\.com/(?:m/space/)?(\d+)(?:/default\.html|\.html)?', "user"),
]


def parse_url(url):
    """从账号链接解析 (platform, pid, kind)。识别失败返回 (None, None, None)。

    kind：user=用户主页，artist/singer=歌手主页。
    """
    url = (url or "").strip()
    if not url:
        return None, None, None
    for plat, pat, kind in _URL_PATTERNS:
        m = re.search(pat, url)
        if m:
            return plat, m.group(1), kind
    # QQ 老式搜索页链接：y.qq.com/portal/search.html?page=...&mid=xxx&type=singer
    if "y.qq.com" in url and re.search(r'type=singer', url):
        m = re.search(r'[?&]mid=([A-Za-z0-9]{10,20})', url)
        if m:
            return "qq", m.group(1), "singer"
    return None, None, None


def extract_pid(platform, url):
    """从链接提取「指定平台」的 (pid, kind)；链接不属于该平台返回 (None, None)。

    手动选平台+粘贴他平台链接时用它校验，避免采集错平台的数据。
    """
    url = (url or "").strip()
    if not url or not platform:
        return None, None
    for plat, pat, kind in _URL_PATTERNS:
        if plat != platform:
            continue
        m = re.search(pat, url)
        if m:
            return m.group(1), kind
    if platform == "qq" and "y.qq.com" in url and re.search(r'type=singer', url):
        m = re.search(r'[?&]mid=([A-Za-z0-9]{10,20})', url)
        if m:
            return m.group(1), "singer"
    return None, None


def collect_account(account):
    """采集单个账号：以账号存储的平台为准提取 ID→调用采集器。

    平台已存但链接解析不出该平台 ID 时报错（如 5sing 账号配了酷我链接），
    绝不能静默改采链接指向的其他平台。
    """
    platform = (account.get("platform") or "").strip()
    pid = (account.get("platform_id") or "").strip()
    kind = (account.get("kind") or "").strip()
    url = account.get("url", "") or ""
    if not pid:
        if platform:
            pid, k = extract_pid(platform, url)
            if not pid:
                return {"platform": platform, "error":
                        "账号链接无法解析出该平台的 ID，请检查链接与所选平台是否匹配"}
            if not kind:
                kind = k
        else:
            platform, pid, kind = parse_url(url)
    if not platform or not pid:
        return {"platform": "unknown", "error": "无法识别链接平台/ID"}

    from .collectors import collect
    return collect(platform, pid, kind)


# ==================== xlsx 长期缓存 ====================

def _xlsx_path(name):
    safe = re.sub(r'[\\/:*?"<>|]', "_", name)
    try:
        from toolbox.config_manager import get_export_subdir
        data_dir = get_export_subdir("music_tracker")
    except Exception:
        data_dir = BASE_DIR / "data"
    data_dir.mkdir(parents=True, exist_ok=True)
    return data_dir / f"{safe}.xlsx"


def save_history(name, record):
    """把采集结果追加到账号的 xlsx 缓存，返回文件路径。"""
    import openpyxl
    path = _xlsx_path(name)
    headers = ["time"] + _RECORD_FIELDS
    row = [datetime.now().strftime("%Y-%m-%d %H:%M:%S")] + [record.get(f, "") for f in _RECORD_FIELDS]
    if path.exists():
        wb = openpyxl.load_workbook(path)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(headers)
    ws.append(row)
    wb.save(path)
    return str(path)


def load_history(name):
    """读取账号历史采集数据（xlsx），返回记录列表（最新在前，每条含 _time）。"""
    import openpyxl
    path = _xlsx_path(name)
    if not path.exists():
        return []
    try:
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception:
        return []
    if len(rows) < 2:
        return []
    headers = rows[0]
    records = []
    for row in rows[1:]:
        if not any(v not in (None, "") for v in row):
            continue
        rec = dict(zip(headers, row))
        time_str = rec.pop("time", "") or ""
        rec["_time"] = time_str
        for k in _NUM_FIELDS:
            if rec.get(k) not in (None, ""):
                try:
                    rec[k] = int(rec[k])
                except (ValueError, TypeError):
                    pass
        records.append(rec)
    records.reverse()  # 最新在前
    return records
