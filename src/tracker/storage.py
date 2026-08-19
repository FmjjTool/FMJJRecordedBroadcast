import json
from datetime import datetime
from pathlib import Path

from .config import OUTPUT_PATHS

try:
    from zoneinfo import ZoneInfo
    _CST = ZoneInfo("Asia/Shanghai")
except Exception:
    _CST = None

_KEY_MAP = {
    "time": "时间",
    "platform": "平台",
    "name": "昵称",
    "fans": "粉丝数",
    "following": "关注数",
    "likes": "获赞数",
    "video_count": "视频数",
    "weibo_count": "微博数",
    "sign": "简介",
    "avatar": "头像",
    "official_title": "认证",
    "level": "等级",
    "vip_type": "会员类型",
    "sex": "性别",
    "verified": "是否认证",
    "verified_reason": "认证说明",
    "gender": "性别",
    "uid": "用户ID",
    "short_id": "短ID",
    "cover": "封面",
    "error": "错误",
    "records": "记录",
}

_REVERSE_KEY_MAP = {v: k for k, v in _KEY_MAP.items()}

_CN_KEYS = [_KEY_MAP[k] for k in (
    "platform", "name", "fans", "following", "likes",
    "video_count", "weibo_count", "sign", "avatar",
    "official_title", "level", "vip_type", "sex",
    "verified", "verified_reason", "gender", "uid",
    "short_id", "cover",
)]

_FLATTEN_KEYS = list(_KEY_MAP.keys())
# remove meta keys
for _skip in ("time", "records", "error"):
    if _skip in _FLATTEN_KEYS:
        _FLATTEN_KEYS.remove(_skip)


def _now_cst() -> str:
    if _CST:
        return datetime.now(_CST).strftime("%Y-%m-%d %H:%M:%S")
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _make_row(time_val: str, record: dict) -> list:
    row = [time_val]
    for k in _FLATTEN_KEYS:
        v = record.get(k)
        row.append(v if v is not None else "")
    return row


def _translate_record(r: dict) -> dict:
    return {_KEY_MAP.get(k, k): v for k, v in r.items()}


def _revert_record(r: dict) -> dict:
    return {_REVERSE_KEY_MAP.get(k, k): v for k, v in r.items()}


def _load(dir_path: Path) -> list:
    fp = dir_path / "records.json"
    if not fp.exists():
        return []
    raw = fp.read_text(encoding="utf-8")
    if not raw.strip():
        return []
    try:
        data = json.loads(raw)
        out = []
        for entry in data:
            t = entry.get(_KEY_MAP["time"], "")
            raw_rec = entry.get(_KEY_MAP["records"], {})
            if isinstance(raw_rec, dict):
                out.append({"time": t, "records": _revert_record(raw_rec)})
        return out
    except (json.JSONDecodeError, KeyError):
        return []


def _write_json(dir_path: Path, history: list):
    cn = []
    for entry in history:
        cn.append({
            _KEY_MAP["time"]: entry["time"],
            _KEY_MAP["records"]: _translate_record(entry["records"]),
        })
    dir_path.mkdir(parents=True, exist_ok=True)
    (dir_path / "records.json").write_text(
        json.dumps(cn, ensure_ascii=False, indent=2), encoding="utf-8",
    )


def _write_xlsx(dir_path: Path, history: list):
    headers = [_KEY_MAP["time"]] + _CN_KEYS
    rows = [headers]
    for idx, entry in enumerate(history):
        if idx > 0:
            rows.append([])
        rows.append(_make_row(entry["time"], entry["records"]))

    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "records"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            ws.cell(row=i + 1, column=j + 1, value=val)
            if i == 0:
                ws.cell(row=i + 1, column=j + 1).font = Font(bold=True)
    col_count = len(headers)
    for ci in range(1, col_count + 1):
        mx = 0
        for ri in range(1, min(len(rows) + 1, 50)):
            v = ws.cell(row=ri, column=ci).value
            if v is not None:
                mx = max(mx, len(str(v)))
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = min(mx + 3, 60)
    wb.save(str(dir_path / "records.xlsx"))


def save(records: list):
    ts = _now_cst()
    for rec in records:
        plat = rec.get("platform")
        if not plat or plat not in OUTPUT_PATHS:
            continue
        dp = OUTPUT_PATHS[plat]
        hist = _load(dp)
        hist.append({"time": ts, "records": rec})
        if len(hist) > 100000:
            hist = hist[-100000:]
        _write_json(dp, hist)
        _write_xlsx(dp, hist)
