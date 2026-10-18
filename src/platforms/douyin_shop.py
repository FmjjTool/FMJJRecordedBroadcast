"""抖音商城（抖音小店）商品信息解析：分享文案 → 商品数据（无需登录）。

链路（2026-10-15 实测）：
  分享文案里的 v.douyin.com 短链
    → 302 跳转到 haohuo.jinritemai.com/ecommerce/trade/detail/index.html
    → Location 查询串里带 id（商品ID）/ sec_author_id / goods_detail（双层 URL 编码 JSON：
      标题、价格区间(分)、销量、主图）
    → detail_schema（双层 URL 编码）里的 meta_params 是 base64(gzip(JSON))，
      解出 shop_id / kol_id / 分享埋点

已能拿到：商品标题 / 商品ID / 价格区间 / 销量 / 主图 / 店铺ID / 带货达人ID / 分享来源
暂拿不到（需登录态或 App）：店铺名称、SKU 规格库存、店铺全部商品列表
——商品页对非 App 打开的分享链只返回「商品预览」数据。

两个坑（已在实现里处理）：
  1. 查询串不能用 urllib.parse.parse_qs：它会把 base64 里的 '+' 转成空格，
     必须手动切分并只做一次 unquote；
  2. detail_schema 是「编码两次」的，要 unquote 两次才能拿到里面的 meta_params。
"""

from __future__ import annotations

import base64
import gzip
import json
import re
import urllib.parse
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path

import httpx

UA_MOBILE = ("Mozilla/5.0 (iPhone; CPU iPhone OS 17_0 like Mac OS X) "
             "AppleWebKit/605.1.15 (KHTML, like Gecko) Version/17.0 "
             "Mobile/15E148 Safari/604.1")

_SHORT_URL_RE = re.compile(r"https?://v\.douyin\.com/[\w\-]+/?")
_ANY_DOUYIN_RE = re.compile(r"https?://[\w.\-]*(?:douyin|iesdouyin)\.com/\S+")

DETAIL_PAGE = "https://haohuo.jinritemai.com/ecommerce/trade/detail/index.html"


# ==================== 数据模型 ====================

@dataclass
class DouyinProduct:
    """抖音商城商品。价格单位：元。"""

    product_id: str = ""
    title: str = ""
    min_price: float = 0.0
    max_price: float = 0.0
    sales: int = 0
    main_image: str = ""
    shop_id: str = ""
    kol_id: str = ""
    sec_author_id: str = ""
    enter_from: str = ""
    app_version: str = ""
    share_url: str = ""
    detail_url: str = ""
    parsed_at: str = ""
    local_dir: str = ""          # 本次下载/导出所在目录（由页面回填）
    image_file: str = ""         # 主图本地文件名（由页面回填）

    @property
    def price_text(self) -> str:
        """价格区间展示文本；最低=最高时只显示一个价格。"""
        if self.min_price <= 0 and self.max_price <= 0:
            return "-"
        if abs(self.max_price - self.min_price) < 1e-9:
            return f"¥{self.min_price:.2f}"
        return f"¥{self.min_price:.2f} ~ ¥{self.max_price:.2f}"

    def to_dict(self) -> dict:
        d = asdict(self)
        d["price_text"] = self.price_text
        return d


# ==================== 查询串 / 参数解码 ====================

def _parse_query_keep_plus(qs: str) -> dict:
    """手动解析查询串：只 unquote，不把 '+' 当空格（否则 base64 会被破坏）。"""
    out = {}
    for part in qs.split("&"):
        if "=" not in part:
            continue
        k, v = part.split("=", 1)
        out[urllib.parse.unquote(k)] = urllib.parse.unquote(v)
    return out


def _decode_goods_detail(location: str) -> dict:
    """goods_detail 是双层 URL 编码的 JSON。"""
    q = _parse_query_keep_plus(urllib.parse.urlparse(location).query)
    raw = q.get("goods_detail")
    if not raw:
        return {}
    try:
        return json.loads(urllib.parse.unquote(urllib.parse.unquote(raw)))
    except Exception:
        return {}


def _decode_detail_schema(location: str) -> tuple[dict, dict]:
    """detail_schema（编码两次）→ (schema 查询串, meta_params 解出的 dict)。"""
    top = _parse_query_keep_plus(urllib.parse.urlparse(location).query)
    schema = top.get("detail_schema")
    if not schema:
        return {}, {}
    schema = urllib.parse.unquote(urllib.parse.unquote(schema))  # 双重编码
    inner = _parse_query_keep_plus(urllib.parse.urlparse(schema).query)
    mp = inner.get("meta_params")
    meta = {}
    if mp:
        try:
            raw = base64.b64decode(mp + "=" * (-len(mp) % 4))
            if raw[:2] == b"\x1f\x8b":            # gzip 魔数
                meta = json.loads(gzip.decompress(raw).decode("utf-8"))
        except Exception:
            meta = {}
    return inner, meta


# ==================== 纯解析（无网络，便于离线测试） ====================

def parse_location(location: str, share_url: str = "") -> DouyinProduct:
    """把 302 跳转后的商品页 URL 解析成 DouyinProduct（不发请求）。"""
    if "haohuo.jinritemai.com" not in location:
        raise ValueError("链接没有跳到抖音商城商品页，可能不是商品分享链")

    q = _parse_query_keep_plus(urllib.parse.urlparse(location).query)
    gd = _decode_goods_detail(location)
    schema_q, meta = _decode_detail_schema(location)

    entrance = {}
    try:
        entrance = json.loads(meta.get("entrance_info") or "{}")
    except Exception:
        entrance = {}

    urls = (gd.get("img") or {}).get("url_list") or []
    product_id = str(q.get("id") or schema_q.get("product_id") or "")

    p = DouyinProduct(
        product_id=product_id,
        title=str(gd.get("title") or ""),
        min_price=float(gd.get("min_price") or 0) / 100.0,
        max_price=float(gd.get("max_price") or 0) / 100.0,
        sales=int(gd.get("sales") or 0),
        main_image=str(urls[0]) if urls else "",
        shop_id=str(entrance.get("shop_id") or schema_q.get("shop_id") or ""),
        kol_id=str(schema_q.get("kol_id") or entrance.get("kol_id") or ""),
        sec_author_id=str(q.get("sec_author_id") or ""),
        enter_from=str(meta.get("enter_from") or ""),
        app_version=str(entrance.get("app_version") or ""),
        share_url=share_url,
        detail_url=f"{DETAIL_PAGE}?id={product_id}" if product_id else location.split("?")[0],
        parsed_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    )
    if not p.product_id and not p.title:
        raise ValueError("链接里没有解析出商品信息（可能链接已失效）")
    return p


def extract_share_url(text: str) -> str:
    """从整段分享文案里抠出抖音链接。"""
    m = _SHORT_URL_RE.search(text or "")
    if m:
        return m.group(0)
    m = _ANY_DOUYIN_RE.search(text or "")
    if m:
        return m.group(0).rstrip("，。!！;；")
    raise ValueError("没找到抖音链接，请把分享文案整段粘贴进来")


# ==================== 网络 ====================

def resolve_short_link(share_url: str, timeout: float = 20.0) -> str:
    """跟随短链的第一跳 302，返回商品页 URL。"""
    with httpx.Client(follow_redirects=False, timeout=timeout) as cli:
        resp = cli.get(share_url, headers={"User-Agent": UA_MOBILE})
        loc = resp.headers.get("location") or ""
        if not loc:
            if resp.status_code == 200 and "haohuo.jinritemai.com" in str(resp.url):
                return str(resp.url)
            raise ValueError(f"短链没有跳转（HTTP {resp.status_code}），可能已失效")
        return loc


def parse_share_text(text: str) -> DouyinProduct:
    """整段分享文案 → DouyinProduct（联网）。"""
    share_url = extract_share_url(text)
    location = resolve_short_link(share_url)
    return parse_location(location, share_url=share_url)


def split_share_texts(text: str) -> list[str]:
    """多行输入 → 每条分享文案（含链接的行才保留）。"""
    out = []
    for line in (text or "").splitlines():
        line = line.strip()
        if line and (".douyin.com" in line or "douyin" in line.lower()):
            out.append(line)
    return out


def parse_many(text: str, progress=None) -> tuple[list[DouyinProduct], list[tuple[str, str]]]:
    """批量解析。返回 (成功列表, [(原文, 错误信息)])。"""
    items = split_share_texts(text)
    ok: list[DouyinProduct] = []
    errs: list[tuple[str, str]] = []
    for i, line in enumerate(items, 1):
        if progress:
            progress(f"解析第 {i}/{len(items)} 条...")
        try:
            ok.append(parse_share_text(line))
        except Exception as e:                       # noqa: BLE001
            errs.append((line[:60], str(e)))
    return ok, errs


# ==================== 本地落盘 ====================

def _safe_name(name: str, max_len: int = 50) -> str:
    """文件名安全化：去掉路径非法字符（与项目其他模块一致的朴素做法）。"""
    s = re.sub(r'[\\/:*?"<>|\r\n\t]', "_", (name or "").strip())
    s = re.sub(r"\s+", " ", s).strip(" .")
    return s[:max_len] or "未命名"


def shop_dir() -> Path:
    """商品数据根目录：download/抖音商品/。"""
    from toolbox.config_manager import get_export_subdir
    return get_export_subdir("douyin_shop")


def product_dir(p: DouyinProduct, create: bool = True) -> Path:
    """单个商品目录：download/抖音商品/<标题>_<商品ID>/。"""
    name = f"{_safe_name(p.title)}_{p.product_id or 'unknown'}"
    d = shop_dir() / name
    if create:
        d.mkdir(parents=True, exist_ok=True)
    return d


def download_main_image(p: DouyinProduct, dest_dir: Path | None = None) -> Path | None:
    """下载商品主图到商品目录，返回文件路径；无图或失败返回 None。"""
    if not p.main_image:
        return None
    d = dest_dir or product_dir(p)
    d.mkdir(parents=True, exist_ok=True)
    suffix = ".png" if ".png" in p.main_image.lower() else ".jpg"
    path = d / f"主图{suffix}"
    try:
        with httpx.Client(timeout=60.0, follow_redirects=True) as cli:
            r = cli.get(p.main_image, headers={"User-Agent": UA_MOBILE})
            r.raise_for_status()
            path.write_bytes(r.content)
        return path
    except Exception:
        return None


def save_product_json(p: DouyinProduct, dest_dir: Path | None = None) -> Path:
    """保存单商品 JSON（含原始字段与展示字段）。"""
    d = dest_dir or product_dir(p)
    d.mkdir(parents=True, exist_ok=True)
    path = d / f"商品信息_{p.product_id or 'unknown'}.json"
    path.write_text(json.dumps(p.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def save_products_json(products: list[DouyinProduct], path: Path | None = None) -> Path:
    """保存整批商品 JSON（默认 download/抖音商品/抖音商品列表_时间戳.json）。"""
    if path is None:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = shop_dir() / f"抖音商品列表_{ts}.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps([p.to_dict() for p in products], ensure_ascii=False, indent=2),
        encoding="utf-8")
    return path


def export_products_xlsx(products: list[DouyinProduct], path: Path | None = None) -> Path:
    """导出商品列表 Excel（商品信息 sheet；本地已有主图的插入图片列）。"""
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if path is None:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = shop_dir() / f"抖音商品列表_{ts}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "商品信息"
    headers = ["序号", "商品标题", "价格区间", "销量", "商品ID", "店铺ID",
               "带货达人ID", "分享来源", "商品页链接", "主图链接", "采集时间"]
    ws.append(headers)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="C2185B")
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [6, 46, 18, 8, 22, 14, 16, 12, 46, 46, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for idx, p in enumerate(products, 1):
        ws.append([idx, p.title, p.price_text, p.sales, p.product_id, p.shop_id,
                   p.kol_id, p.enter_from, p.detail_url, p.main_image, p.parsed_at])
        # 本地已有主图 → 插入到该行（缩放到 ~120px 高）
        if p.local_dir and p.image_file:
            img_path = Path(p.local_dir) / p.image_file
            if img_path.exists():
                try:
                    img = XLImage(str(img_path))
                    ratio = 120.0 / (img.height or 120)
                    img.height = int(img.height * ratio)
                    img.width = int(img.width * ratio)
                    ws.row_dimensions[idx + 1].height = img.height
                    ws.add_image(img, f"L{idx + 1}")
                except Exception:
                    pass
    if products and any(p.image_file for p in products):
        ws.column_dimensions["L"].width = 20

    wb.save(str(path))
    return path
