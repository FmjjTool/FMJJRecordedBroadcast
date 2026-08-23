"""一次性迁移脚本：把老项目（平台信息获取）download 数据合并到本项目 download/。

规则：
- 普通文件：复制（同路径已存在则跳过，新项目数据优先），老项目原文件不动。
- danmaku/、stats/ 旧目录 → 并入 弹幕文件/（新项目两者同目录）。
- 账号追踪数据/*.xlsx：按行合并（老行+新行，整行去重，按 time 排序），合并前备份新文件。
- 已获取数据：保持原结构复制（根目录文件靠文件名平台后缀识别；bilibili/douyin/weibo
  英文子目录靠目录名识别；B站/抖音/微博中文子目录只放媒体，资产管理页递归扫描可识别）。

用法：python migrate_old_data.py [--dry]
"""

import shutil
import sys
from pathlib import Path

OLD = Path(r"D:\PythonApps\all\平台信息获取\download")
NEW = Path(r"D:\PythonApps\all\PySide6\download")

copied, skipped, merged_rows, errors = [], [], [], []


def copy_dir(src: Path, dst: Path, label: str):
    if not src.exists():
        print(f"[跳过] {label}: 源目录不存在")
        return
    for f in sorted(src.rglob("*")):
        if not f.is_file():
            continue
        rel = f.relative_to(src)
        target = dst / rel
        if target.exists():
            skipped.append(f"{label}/{rel}")
            continue
        try:
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(f, target)
            copied.append(f"{label}/{rel}")
        except Exception as e:
            errors.append(f"{label}/{rel}: {e}")


def merge_tracker_xlsx():
    """账号追踪数据：png 等普通文件直接复制；同名 xlsx 老行并进新文件（整行去重、按 time 升序），合并前备份新文件。"""
    import openpyxl

    old_dir = OLD / "账号追踪数据"
    new_dir = NEW / "账号追踪数据"
    if not old_dir.exists():
        return
    for f in sorted(old_dir.rglob("*")):
        if f.is_file() and f.suffix.lower() != ".xlsx":
            target = new_dir / f.relative_to(old_dir)
            if target.exists():
                skipped.append(f"账号追踪数据/{f.relative_to(old_dir)}")
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(f, target)
            copied.append(f"账号追踪数据/{f.relative_to(old_dir)}")
    for ox in sorted(old_dir.glob("*.xlsx")):
        nx = new_dir / ox.name
        if not nx.exists():
            target = new_dir / ox.name
            shutil.copy2(ox, target)
            copied.append(f"账号追踪数据/{ox.name}")
            continue
        owb = openpyxl.load_workbook(ox)
        nwb = openpyxl.load_workbook(nx)
        ows, nws = owb.active, nwb.active
        orows = [tuple(r) for r in ows.iter_rows(values_only=True)]
        nrows = [tuple(r) for r in nws.iter_rows(values_only=True)]
        if not orows or not nrows or orows[0] != nrows[0]:
            print(f"[跳过] {ox.name}: 表头不一致，不合并")
            continue
        header, odata, ndata = orows[0], orows[1:], nrows[1:]
        seen = set(ndata)
        add = [r for r in odata if r not in seen]
        if not add:
            skipped.append(f"账号追踪数据/{ox.name}（老数据已全部包含）")
            continue
        bak = nx.with_suffix(".xlsx.bak")
        if not bak.exists():
            shutil.copy2(nx, bak)
        all_rows = ndata + add
        t_idx = header.index("time") if "time" in header else 0
        all_rows.sort(key=lambda r: str(r[t_idx] or ""))
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(list(header))
        for r in all_rows:
            ws.append(list(r))
        wb.save(nx)
        merged_rows.append(f"账号追踪数据/{ox.name}: 追加 {len(add)} 行（老 {len(odata)} + 新 {len(ndata)}）")


def main():
    dry = "--dry" in sys.argv
    print(f"老目录: {OLD}\n新目录: {NEW}\n模式: {'预演(不写入)' if dry else '复制合并'}\n")

    plan = [
        (OLD / "已获取数据", NEW / "已获取数据", "已获取数据"),
        (OLD / "弹幕文件", NEW / "弹幕文件", "弹幕文件"),
        (OLD / "danmaku", NEW / "弹幕文件", "danmaku→弹幕文件"),
        (OLD / "stats", NEW / "弹幕文件", "stats→弹幕文件"),
        (OLD / "字幕文件", NEW / "字幕文件", "字幕文件"),
        (OLD / "歌曲文件", NEW / "歌曲文件", "歌曲文件"),
        (OLD / "导出的图片", NEW / "导出的图片", "导出的图片"),
        (OLD / "音乐信息", NEW / "音乐信息", "音乐信息"),
    ]
    if dry:
        for src, dst, label in plan:
            if not src.exists():
                print(f"[跳过] {label}: 源目录不存在")
                continue
            n_new = sum(1 for f in src.rglob("*") if f.is_file() and not (dst / f.relative_to(src)).exists())
            n_skip = sum(1 for f in src.rglob("*") if f.is_file() and (dst / f.relative_to(src)).exists())
            print(f"[预演] {label}: 新复制 {n_new} 个，已存在跳过 {n_skip} 个")
        print("[预演] 账号追踪数据 xlsx 将按行合并（老行追加，time 排序）")
        return

    for src, dst, label in plan:
        copy_dir(src, dst, label)
    merge_tracker_xlsx()

    print(f"\n===== 完成 =====")
    print(f"新复制: {len(copied)} 个文件")
    for p in copied:
        print(f"  + {p}")
    if merged_rows:
        print(f"\nxlsx 合并:")
        for m in merged_rows:
            print(f"  ~ {m}")
    if skipped:
        print(f"\n已存在跳过: {len(skipped)} 个（新项目数据优先，不覆盖）")
        for p in skipped:
            print(f"  = {p}")
    if errors:
        print(f"\n失败: {len(errors)} 个")
        for e in errors:
            print(f"  ! {e}")


if __name__ == "__main__":
    main()
